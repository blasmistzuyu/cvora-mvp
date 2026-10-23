import io
from pathlib import Path
from fastapi.testclient import TestClient
from docx import Document
from pypdf import PdfReader
from main import app, demo_analysis

client = TestClient(app)
CV = 'Nombre Apellido\ncorreo@ejemplo.com\nEDUCACIÓN\nSecundario completo\nPROYECTOS\nParticipé en la organización de una feria estudiantil.\nHABILIDADES\nExcel básico y comunicación.'


def test_health_and_home():
    assert client.get('/api/health').status_code == 200
    page = client.get('/')
    assert page.status_code == 200 and 'CVORA' in page.text
    assert client.get('/static/styles.css').status_code == 200
    assert client.get('/privacidad').status_code == 200


def test_analyze_demo(monkeypatch):
    monkeypatch.delenv('OPENAI_API_KEY', raising=False)
    r = client.post('/api/analyze', json={'cv': CV, 'job': 'Atención al cliente con Excel y SAP', 'profile': 'primer-empleo'})
    assert r.status_code == 200
    assert r.json()['mode'] == 'demo'
    assert 'sap' in r.json()['keywords_to_review']
    assert r.json()['rewritten_cv']
    assert 'inventes' not in r.json()['rewritten_cv'].lower()


def test_missing_fields_and_bad_uploads():
    assert client.post('/api/analyze',json={'cv':'corto'}).status_code == 422
    assert client.post('/api/extract',files={'file': ('hello.exe', b'xxx', 'application/octet-stream')}).status_code == 415
    assert client.post('/api/extract',files={'file': ('big.txt', b'A' * (4*1024*1024+1), 'text/plain')}).status_code == 413


def test_text_and_word_extraction():
    txt = client.post('/api/extract', files={'file': ('cv.txt', CV.encode(), 'text/plain')})
    assert txt.status_code == 200 and 'Nombre' in txt.json()['text']
    doc=Document();doc.add_paragraph(CV)
    b=io.BytesIO();doc.save(b)
    word = client.post('/api/extract', files={'file': ('cv.docx', b.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
    assert word.status_code == 200 and 'Excel' in word.json()['text']


def test_pdf_docx_exports():
    docx_res=client.post('/api/export/docx',json={'content':CV})
    assert docx_res.status_code == 200
    doc = Document(io.BytesIO(docx_res.content))
    assert any('Nombre' in p.text for p in doc.paragraphs)
    pdf_res=client.post('/api/export/pdf',json={'content':CV})
    assert pdf_res.status_code == 200
    pdf = PdfReader(io.BytesIO(pdf_res.content))
    assert len(pdf.pages)>=1
    assert 'Nombre' in pdf.pages[0].extract_text()


def test_escaped_strings_in_pdf():
    pdf_res=client.post('/api/export/pdf',json={'content':'Nombre <Apellido> & Empresa\nEXPERIENCIA\n• A & B'})
    assert pdf_res.status_code == 200


def test_pdf_extraction():
    from main import export_pdf, ExportPayload
    response=export_pdf(ExportPayload(content=CV))
    parsed=client.post('/api/extract',files={'file': ('sample.pdf',response.body,'application/pdf')})
    assert parsed.status_code == 200 and 'Nombre' in parsed.json()['text']

def test_ai_requires_opt_in(monkeypatch):
    monkeypatch.setenv('OPENAI_API_KEY','placeholder-key')
    monkeypatch.setenv('CVORA_ENABLE_AI','true')
    from main import app
    import main
    monkeypatch.setattr(main,'ai_analysis',lambda cv,job,profile: {'mode':'ai','headline':'Listo','suggestions':[],'keywords_to_review':[],'rewritten_cv':cv,'notice':'Revisar'})
    without=client.post('/api/analyze',json={'cv':CV,'ai_consent':False})
    assert without.status_code == 200 and without.json()['mode']=='demo'
    yes=client.post('/api/analyze',json={'cv':CV,'ai_consent':True})
    assert yes.status_code == 200 and yes.json()['mode']=='ai'
